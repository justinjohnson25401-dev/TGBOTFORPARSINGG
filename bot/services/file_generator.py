"""
Excel file service for pack files management
"""
import os
import io
import logging
from typing import Optional, Tuple, List
import openpyxl
from openpyxl import Workbook

from bot.config import TEMP_FILES_DIR, DEMO_FILE_PATH, GDRIVE_BASE_FOLDER_ID, GDRIVE_DEMO_FOLDER_ID, CONTACTS_PER_PACK
from bot.services.gdrive import get_file_as_bytes, ensure_temp_dir, list_files_in_folder, download_file

logger = logging.getLogger(__name__)


def get_pack_filename(city: str, category: str, pack_number: int) -> str:
    """
    Generate pack filename based on naming convention

    Format: {City}_{Category}_PACK_{NN}.xlsx
    Example: Moscow_SalonKrasoty_PACK_01.xlsx
    """
    return f"{city}_{category}_PACK_{pack_number:02d}.xlsx"


def parse_pack_filename(filename: str) -> Optional[dict]:
    """
    Parse pack filename to extract city, category and pack number

    Returns dict with keys: city, category, pack_number or None if invalid
    """
    try:
        # Remove extension
        name = filename.replace('.xlsx', '').replace('.XLSX', '')
        parts = name.split('_')

        if len(parts) < 4:
            return None

        city = parts[0]
        category = parts[1]

        # Find PACK part
        pack_idx = None
        for i, part in enumerate(parts):
            if part.upper() == 'PACK':
                pack_idx = i
                break

        if pack_idx is None or pack_idx + 1 >= len(parts):
            return None

        pack_number = int(parts[pack_idx + 1])

        return {
            'city': city,
            'category': category,
            'pack_number': pack_number
        }
    except Exception as e:
        logger.error(f"Error parsing filename {filename}: {e}")
        return None


async def get_available_packs(city: str, category: str) -> List[dict]:
    """
    Get list of available pack files for city+category from Google Drive

    Returns list of dicts with keys: pack_number, file_id, filename
    """
    try:
        files = await list_files_in_folder(GDRIVE_BASE_FOLDER_ID)

        packs = []
        prefix = f"{city}_{category}_PACK_"

        for file in files:
            filename = file.get('name', '')
            if filename.startswith(prefix) and filename.endswith('.xlsx'):
                parsed = parse_pack_filename(filename)
                if parsed:
                    packs.append({
                        'pack_number': parsed['pack_number'],
                        'file_id': file.get('id'),
                        'filename': filename
                    })

        # Sort by pack number
        packs.sort(key=lambda x: x['pack_number'])
        return packs

    except Exception as e:
        logger.error(f"Error getting available packs: {e}")
        return []


async def download_pack_file(city: str, category: str, pack_number: int) -> Optional[Tuple[str, bytes]]:
    """
    Download specific pack file from Google Drive

    Args:
        city: City code (e.g., "Moscow")
        category: Category code (e.g., "SalonKrasoty")
        pack_number: Pack number (1, 2, 3, etc.)

    Returns:
        Tuple of (filename, file_bytes) or None if not found
    """
    try:
        filename = get_pack_filename(city, category, pack_number)
        logger.info(f"Downloading pack file: {filename}")

        # Find file in Google Drive
        files = await list_files_in_folder(GDRIVE_BASE_FOLDER_ID)

        for file in files:
            if file.get('name', '').upper() == filename.upper():
                file_id = file.get('id')
                file_bytes = await get_file_as_bytes(file_id)

                if file_bytes:
                    return filename, file_bytes
                else:
                    logger.error(f"Failed to download file: {filename}")
                    return None

        logger.error(f"Pack file not found: {filename}")
        return None

    except Exception as e:
        logger.error(f"Error downloading pack file: {e}")
        return None


async def get_pack_files_for_order(
    city: str,
    category: str,
    pack_count: int,
    already_purchased_packs: List[int] = None
) -> List[Tuple[str, bytes]]:
    """
    Get all pack files for an order

    Args:
        city: City code
        category: Category code
        pack_count: Number of packs to get (1 pack = 1000 contacts)
        already_purchased_packs: List of pack numbers user already has

    Returns:
        List of (filename, file_bytes) tuples
    """
    if already_purchased_packs is None:
        already_purchased_packs = []

    files = []
    available_packs = await get_available_packs(city, category)

    # Find next available packs that user doesn't have
    packs_to_download = []
    for pack in available_packs:
        if pack['pack_number'] not in already_purchased_packs:
            packs_to_download.append(pack['pack_number'])
            if len(packs_to_download) >= pack_count:
                break

    # Download each pack
    for pack_num in packs_to_download:
        result = await download_pack_file(city, category, pack_num)
        if result:
            files.append(result)

    return files


async def get_demo_file_from_gdrive() -> Optional[Tuple[str, bytes]]:
    """
    Get demo file from Google Drive demo folder

    Returns:
        Tuple of (filename, file_bytes) or None
    """
    try:
        if not GDRIVE_DEMO_FOLDER_ID:
            logger.warning("GDRIVE_DEMO_FOLDER_ID not configured")
            return get_demo_file()

        files = await list_files_in_folder(GDRIVE_DEMO_FOLDER_ID)

        # Find DEMO.xlsx or any xlsx file
        for file in files:
            filename = file.get('name', '')
            if filename.upper().endswith('.XLSX'):
                file_id = file.get('id')
                file_bytes = await get_file_as_bytes(file_id)

                if file_bytes:
                    return "DEMO_Base.xlsx", file_bytes

        logger.warning("No demo file found in Google Drive")
        return get_demo_file()

    except Exception as e:
        logger.error(f"Error getting demo file from GDrive: {e}")
        return get_demo_file()


def get_demo_file() -> Optional[Tuple[str, bytes]]:
    """
    Get demo file for preview (local fallback)

    Returns:
        Tuple of (filename, file_bytes) or None
    """
    try:
        if os.path.exists(DEMO_FILE_PATH):
            with open(DEMO_FILE_PATH, 'rb') as f:
                return "DEMO_Base.xlsx", f.read()

        # Generate a demo file if not exists
        return generate_demo_file()

    except Exception as e:
        logger.error(f"Error getting demo file: {e}")
        return None


def generate_demo_file() -> Optional[Tuple[str, bytes]]:
    """
    Generate a demo file with masked data

    Returns:
        Tuple of (filename, file_bytes) or None
    """
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Demo"

        # Headers
        headers = [
            "Название", "Телефон", "Telegram", "VK", "WhatsApp",
            "Email", "Сайт", "Адрес", "Рейтинг", "Отзывы"
        ]
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)

        # Demo data (masked)
        demo_data = [
            ["Салон красоты 'Элегант'", "+79999999999", "@hidden", "vk.com/hidden", "+79999999999",
             "hidden@email.com", "example.com", "г. Москва, ул. Примерная, 1", "4.8", "150"],
            ["Beauty Studio", "+79999999999", "@hidden", "vk.com/hidden", "+79999999999",
             "hidden@email.com", "example.com", "г. Москва, ул. Примерная, 2", "4.5", "89"],
            ["Студия 'Красота'", "+79999999999", "@hidden", "vk.com/hidden", "+79999999999",
             "hidden@email.com", "example.com", "г. Москва, ул. Примерная, 3", "4.9", "234"],
            ["Nail Art Studio", "+79999999999", "@hidden", "vk.com/hidden", "+79999999999",
             "hidden@email.com", "example.com", "г. Москва, ул. Примерная, 4", "4.7", "112"],
            ["Парикмахерская 'Стиль'", "+79999999999", "@hidden", "vk.com/hidden", "+79999999999",
             "hidden@email.com", "example.com", "г. Москва, ул. Примерная, 5", "4.6", "78"],
        ]

        for row_idx, row_data in enumerate(demo_data, 2):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        # Adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[column].width = min(max_length + 2, 30)

        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        # Also save to disk
        ensure_temp_dir()
        os.makedirs(os.path.dirname(DEMO_FILE_PATH), exist_ok=True)
        with open(DEMO_FILE_PATH, 'wb') as f:
            f.write(output.getvalue())

        output.seek(0)
        return "DEMO_Base.xlsx", output.read()

    except Exception as e:
        logger.error(f"Error generating demo file: {e}")
        return None


def get_pack_count_for_contacts(contacts: int) -> int:
    """Calculate number of packs needed for given contact count"""
    return contacts // CONTACTS_PER_PACK


def get_contacts_for_packs(pack_count: int) -> int:
    """Calculate total contacts for given pack count"""
    return pack_count * CONTACTS_PER_PACK


# Keep old function for backward compatibility
async def generate_pack_file(
        gdrive_file_id: str,
        pack_number: int,
        pack_size: int,
        city: str,
        category: str
) -> Optional[Tuple[str, bytes]]:
    """
    Legacy function - now just downloads pre-made pack file
    """
    return await download_pack_file(city, category, pack_number)
